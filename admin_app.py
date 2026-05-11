#!/usr/bin/env python3
"""Local STOURS catalogue admin app.

Run from the repository root:

    python3 admin_app.py

Then open http://127.0.0.1:8000/admin
"""

from __future__ import annotations

import argparse
import json
import mimetypes
import re
import shutil
import time
import unicodedata
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import quote_plus, unquote, urlparse

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parent
DEFAULT_CATALOGUE = ROOT / "index.html"
BACKUP_DIR = ROOT / ".admin_backups"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
SECTION_LABELS = ["ROOMS", "RESTAURANTS", "SIGNATURE EXPERIENCE", "FACILITIES"]


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", normalized.lower()).strip("-")
    return slug or "hotel"


def normalize(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", " ", normalized.upper()).strip()


def read_soup(catalogue_path: Path) -> BeautifulSoup:
    return BeautifulSoup(catalogue_path.read_text(encoding="utf-8"), "html.parser")


def make_backup(catalogue_path: Path) -> Path:
    BACKUP_DIR.mkdir(exist_ok=True)
    backup = BACKUP_DIR / f"{catalogue_path.stem}-{time.strftime('%Y%m%d-%H%M%S')}.html"
    shutil.copy2(catalogue_path, backup)
    return backup


def write_soup(catalogue_path: Path, soup: BeautifulSoup) -> Path:
    backup = make_backup(catalogue_path)
    catalogue_path.write_text(str(soup), encoding="utf-8")
    return backup


def text_or_empty(node: Tag | None) -> str:
    return node.get_text(" ", strip=True) if node else ""


def hotel_cards(soup: BeautifulSoup) -> list[Tag]:
    return list(soup.select("article.hotel-card"))


def parent_city(card: Tag) -> dict[str, str]:
    section = card.find_parent("section")
    return {
        "id": section.get("id", "") if section else "",
        "name": text_or_empty(section.select_one(".city-title, h2")) if section else "",
    }


def card_images(card: Tag) -> list[str]:
    hero = card.select_one(".visual > img")
    images = [hero.get("src", "")] if hero and hero.get("src") else []
    for img in card.select(".thumbs img"):
        src = img.get("src", "")
        if src:
            images.append(src)
    return images


def card_to_dict(card: Tag, index: int) -> dict[str, object]:
    city = parent_city(card)
    data = {
        title.get_text(strip=True): text_or_empty(title.find_next_sibling(class_="data-text"))
        for title in card.select(".data-title")
    }
    details = {
        text_or_empty(detail.select_one("small")): text_or_empty(detail.select_one("span"))
        for detail in card.select(".travel-detail")
    }
    badges = [text_or_empty(pill) for pill in card.select(".pill")]
    return {
        "index": index,
        "city": city,
        "name": text_or_empty(card.select_one("h3")),
        "stars": text_or_empty(card.select_one(".stars")),
        "badges": badges,
        "classification": badges[0] if badges else "",
        "location": details.get("Localisation Google Maps", city["name"]),
        "airport_route": details.get("Trajet hôtel ↔ aéroport", ""),
        "room_names": details.get("Chambres", data.get("ROOMS", "")),
        "official_url": card.select_one(".hotel-link").get("href", "") if card.select_one(".hotel-link") else "",
        "sections": {label: data.get(label, "") for label in SECTION_LABELS},
        "images": card_images(card),
    }


def all_hotels(catalogue_path: Path) -> list[dict[str, object]]:
    soup = read_soup(catalogue_path)
    return [card_to_dict(card, index) for index, card in enumerate(hotel_cards(soup))]


def set_text(node: Tag | None, value: str) -> None:
    if node is not None:
        node.clear()
        node.append(value)


def google_maps_search(name: str, destination: str) -> str:
    query = f"{name} {destination} Morocco hotel".strip()
    return "https://www.google.com/maps/search/?api=1&query=" + quote_plus(query)


def google_maps_directions(name: str, destination: str) -> str:
    origin = f"{name} {destination} Morocco".strip()
    airport = f"{destination} airport Morocco".strip()
    return (
        "https://www.google.com/maps/dir/?api=1&origin="
        + quote_plus(origin)
        + "&destination="
        + quote_plus(airport)
        + "&travelmode=driving"
    )


def ensure_travel_details(card: Tag, soup: BeautifulSoup) -> Tag:
    details = card.select_one(".travel-details")
    if details:
        return details
    body = card.select_one(".hotel-body")
    if body is None:
        raise ValueError("Hotel card has no body")
    details = soup.new_tag("div")
    details["class"] = "travel-details"
    name = text_or_empty(card.select_one("h3"))
    city = parent_city(card)["name"]
    rooms = ""
    for data_section in card.select(".data-section"):
        if text_or_empty(data_section.select_one(".data-title")) == "ROOMS":
            rooms = text_or_empty(data_section.select_one(".data-text"))
    for label, value, href, cta in [
        ("Localisation Google Maps", city, google_maps_search(name, city), "Voir la carte"),
        ("Trajet hôtel ↔ aéroport", "Trajet à vérifier sur Google Maps", google_maps_directions(name, city), "Ouvrir le trajet"),
        ("Chambres", rooms or "Noms/types de chambres à compléter depuis la source officielle", None, None),
    ]:
        item = soup.new_tag("div")
        item["class"] = "travel-detail"
        small = soup.new_tag("small")
        small.append(label)
        span = soup.new_tag("span")
        span.append(value)
        item.append(small)
        item.append(span)
        if href:
            link = soup.new_tag("a", href=href, target="_blank", rel="noopener")
            link.append(cta or "Ouvrir")
            item.append(link)
        details.append(item)
    thumbs = card.select_one(".thumbs")
    if thumbs:
        thumbs.insert_before(details)
    else:
        body.append(details)
    return details


def set_travel_detail(card: Tag, soup: BeautifulSoup, label: str, value: str, href: str | None = None) -> None:
    details = ensure_travel_details(card, soup)
    item = None
    for candidate in details.select(".travel-detail"):
        if text_or_empty(candidate.select_one("small")) == label:
            item = candidate
            break
    if item is None:
        item = soup.new_tag("div")
        item["class"] = "travel-detail"
        small = soup.new_tag("small")
        small.append(label)
        span = soup.new_tag("span")
        item.append(small)
        item.append(span)
        details.append(item)
    set_text(item.select_one("span"), value)
    if href:
        link = item.select_one("a")
        if link is None:
            link = soup.new_tag("a", target="_blank", rel="noopener")
            link.append("Ouvrir")
            item.append(link)
        link["href"] = href


def find_card(soup: BeautifulSoup, index: int) -> Tag:
    cards = hotel_cards(soup)
    if index < 0 or index >= len(cards):
        raise ValueError(f"Hotel index {index} is out of range")
    return cards[index]


def set_card_images(card: Tag, images: list[str]) -> None:
    clean_images = [image for image in images if image]
    if not clean_images:
        return
    name = text_or_empty(card.select_one("h3"))
    hero = card.select_one(".visual > img")
    if hero:
        hero["src"] = clean_images[0]
        hero["alt"] = name
    thumbs = card.select_one(".thumbs")
    if thumbs is None:
        body = card.select_one(".hotel-body")
        thumbs = card.new_tag("div") if hasattr(card, "new_tag") else None
        if body and thumbs:
            thumbs["class"] = "thumbs"
            body.append(thumbs)
    if thumbs is None:
        return
    thumbs.clear()
    soup = card if isinstance(card, BeautifulSoup) else card.find_parent()
    while soup and not isinstance(soup, BeautifulSoup):
        soup = soup.find_parent()
    if soup is None:
        return
    for src in clean_images[1:]:
        button = soup.new_tag("button", type="button")
        button["onclick"] = f"openLightbox('{src}')"
        image = soup.new_tag("img", alt=f"{name} visual", loading="lazy", src=src)
        button.append(image)
        thumbs.append(button)


def update_hotel(catalogue_path: Path, index: int, payload: dict[str, object]) -> dict[str, object]:
    soup = read_soup(catalogue_path)
    card = find_card(soup, index)
    name = str(payload.get("name", "")).strip()
    if name:
        set_text(card.select_one("h3"), name)
        hero = card.select_one(".visual > img")
        if hero:
            hero["alt"] = name
        for img in card.select(".thumbs img"):
            img["alt"] = f"{name} visual"

    sections = payload.get("sections")
    if isinstance(sections, dict):
        for data_section in card.select(".data-section"):
            label = text_or_empty(data_section.select_one(".data-title"))
            if label in sections:
                set_text(data_section.select_one(".data-text"), str(sections[label]))

    if "classification" in payload:
        first_pill = card.select_one(".pill")
        if first_pill:
            set_text(first_pill, str(payload["classification"]))
    hotel_name = text_or_empty(card.select_one("h3"))
    destination = str(payload.get("location", parent_city(card)["name"]))
    if "location" in payload:
        set_travel_detail(card, soup, "Localisation Google Maps", destination, google_maps_search(hotel_name, destination))
    if "airport_route" in payload:
        set_travel_detail(
            card,
            soup,
            "Trajet hôtel ↔ aéroport",
            str(payload["airport_route"]),
            google_maps_directions(hotel_name, destination),
        )
    if "room_names" in payload:
        set_travel_detail(card, soup, "Chambres", str(payload["room_names"]))
    if "official_url" in payload:
        link = card.select_one(".hotel-link")
        if link:
            link["href"] = str(payload["official_url"])
    if isinstance(payload.get("images"), list):
        set_card_images(card, [str(image) for image in payload["images"]])

    backup = write_soup(catalogue_path, soup)
    return {"hotel": card_to_dict(find_card(read_soup(catalogue_path), index), index), "backup": str(backup)}


def validate_image(path: Path) -> tuple[int, int]:
    with Image.open(path) as image:
        image.verify()
    with Image.open(path) as image:
        return image.size


def download_photo(catalogue_path: Path, index: int, url: str, slot: int) -> dict[str, object]:
    if not url.lower().startswith(("http://", "https://")):
        raise ValueError("Photo URL must start with http:// or https://")
    soup = read_soup(catalogue_path)
    card = find_card(soup, index)
    hotel_name = text_or_empty(card.select_one("h3"))
    hotel_slug = slugify(hotel_name)
    hotel_dir = ROOT / "assets" / "official" / hotel_slug
    hotel_dir.mkdir(parents=True, exist_ok=True)

    response = requests.get(url, timeout=30, headers={"User-Agent": "STOURS local catalogue admin/1.0"})
    response.raise_for_status()
    content_type = response.headers.get("content-type", "").split(";")[0].strip()
    ext = mimetypes.guess_extension(content_type) or Path(urlparse(url).path).suffix or ".jpg"
    ext = ".jpg" if ext == ".jpe" else ext.lower()
    if ext not in IMAGE_EXTENSIONS:
        ext = ".jpg"
    filename = f"{slot + 1:02d}{ext}"
    target = hotel_dir / filename
    target.write_bytes(response.content)
    width, height = validate_image(target)

    relative_path = target.relative_to(ROOT).as_posix()
    images = card_images(card)
    while len(images) <= slot:
        images.append("")
    images[slot] = relative_path
    set_card_images(card, images)
    backup = write_soup(catalogue_path, soup)
    return {
        "path": relative_path,
        "width": width,
        "height": height,
        "hotel": card_to_dict(find_card(read_soup(catalogue_path), index), index),
        "backup": str(backup),
    }


def parse_excel(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    header_row_number = None
    headers: list[str] = []
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        values = [str(value).strip() if value is not None else "" for value in row]
        normalized = [normalize(value) for value in values]
        if any(value in {"HOTEL NAME", "HOTEL", "NOM HOTEL", "NOM DE L HOTEL"} for value in normalized):
            header_row_number = row_number
            headers = normalized
            break
    if header_row_number is None:
        raise ValueError("Could not find an Excel header row with a hotel-name column")

    def find_column(candidates: set[str]) -> int | None:
        for column_index, header in enumerate(headers):
            if header in candidates:
                return column_index
        return None

    city_col = find_column({"CITY", "VILLE", "DESTINATION"})
    hotel_col = find_column({"HOTEL NAME", "HOTEL", "NOM HOTEL", "NOM DE L HOTEL"})
    classification_col = find_column({"CLASSIFICATION", "CATEGORY", "CATEGORIE", "CATEGORIE HOTEL"})
    link_col = find_column({"LINK", "URL", "WEBSITE", "SITE WEB", "LIEN"})
    if hotel_col is None:
        raise ValueError("Could not identify the hotel-name column")
    if classification_col is None and hotel_col + 1 < len(headers):
        classification_col = hotel_col + 1
    if link_col is None and hotel_col + 2 < len(headers):
        link_col = hotel_col + 2

    rows: list[dict[str, str]] = []
    current_city = ""
    for row in worksheet.iter_rows(min_row=header_row_number + 1, values_only=True):
        values = [str(value).strip() if value is not None else "" for value in row]
        if city_col is not None and city_col < len(values) and values[city_col]:
            current_city = values[city_col]
        hotel_name = values[hotel_col] if hotel_col < len(values) else ""
        if not hotel_name:
            continue
        rows.append(
            {
                "city": current_city,
                "name": hotel_name,
                "classification": values[classification_col] if classification_col is not None and classification_col < len(values) else "",
                "official_url": values[link_col] if link_col is not None and link_col < len(values) else "",
            }
        )
    return rows


def clone_card_template(soup: BeautifulSoup, row: dict[str, str]) -> Tag:
    source = soup.select_one("article.hotel-card")
    if source is None:
        raise ValueError("No hotel card template found")
    card = BeautifulSoup(str(source), "html.parser").select_one("article.hotel-card")
    if card is None:
        raise ValueError("Could not clone hotel-card template")
    name = row["name"]
    city = row.get("city", "")
    classification = row.get("classification", "") or "À compléter"
    set_text(card.select_one("h3"), name)
    hero = card.select_one(".visual > img")
    if hero:
        hero["src"] = "assets/logos/stours-dmc-preferred.png"
        hero["alt"] = name
    for node in card.select(".pill, .badge-row, .stars"):
        node.decompose()
    for data_section in card.select(".data-section"):
        label = text_or_empty(data_section.select_one(".data-title"))
        if label == "ROOMS":
            value = f"Informations chambres à compléter pour {name}."
        elif label == "RESTAURANTS":
            value = f"Informations restauration à compléter pour {name}."
        elif label == "SIGNATURE EXPERIENCE":
            value = f"Expérience signature à compléter pour {name}."
        else:
            value = "Localisation et trajet consultables via Google Maps."
        set_text(data_section.select_one(".data-text"), value)
    for node in card.select(".info-grid, .photo-note, .travel-details"):
        node.decompose()
    set_travel_detail(card, soup, "Localisation Google Maps", city, google_maps_search(name, city))
    set_travel_detail(
        card,
        soup,
        "Trajet hôtel ↔ aéroport",
        "Trajet à vérifier sur Google Maps",
        google_maps_directions(name, city),
    )
    set_travel_detail(card, soup, "Chambres", f"Noms/types de chambres à compléter pour {name}.")
    thumbs = card.select_one(".thumbs")
    if thumbs:
        thumbs.clear()
    link = card.select_one(".hotel-link")
    if link:
        link["href"] = row.get("official_url", "")
    for source in card.select(".source"):
        source.decompose()
    return card


def recalculate_counts(soup: BeautifulSoup) -> None:
    for section in soup.select("section.section"):
        grid = section.select_one(".hotels-grid")
        counter = section.select_one(".city-count b")
        if grid and counter:
            set_text(counter, str(len(grid.select("article.hotel-card"))))
    hero_metrics = soup.select(".metric strong")
    total = len(hotel_cards(soup))
    if hero_metrics:
        set_text(hero_metrics[0], str(total))


def apply_excel_import(catalogue_path: Path, excel_path: Path) -> dict[str, object]:
    rows = parse_excel(excel_path)
    soup = read_soup(catalogue_path)
    cards = hotel_cards(soup)
    by_city_name: dict[tuple[str, str], Tag] = {}
    by_name: dict[str, list[Tag]] = {}
    for card in cards:
        name_key = normalize(text_or_empty(card.select_one("h3")))
        by_name.setdefault(name_key, []).append(card)
        city_values = {normalize(parent_city(card)["name"])}
        location = card_to_dict(card, 0).get("location")
        if location:
            city_values.add(normalize(str(location)))
        for city_value in city_values:
            by_city_name[(city_value, name_key)] = card
    sections = {normalize(text_or_empty(section.select_one(".city-title, h2"))): section for section in soup.select("section.section")}
    updated = 0
    created = 0
    missing_city_sections: list[str] = []

    for row in rows:
        key = normalize(row["name"])
        city_key = normalize(row.get("city", ""))
        card = by_city_name.get((city_key, key))
        if card is None and not city_key and len(by_name.get(key, [])) == 1:
            card = by_name[key][0]
        if card is None:
            section = sections.get(city_key)
            grid = section.select_one(".hotels-grid") if section else None
            if grid is None:
                missing_city_sections.append(row.get("city", ""))
                continue
            card = clone_card_template(soup, row)
            grid.append(card)
            by_name.setdefault(key, []).append(card)
            by_city_name[(city_key, key)] = card
            created += 1
        else:
            updated += 1
        set_text(card.select_one("h3"), row["name"])
        if row.get("city"):
            set_travel_detail(card, soup, "Localisation Google Maps", row["city"].strip(), google_maps_search(row["name"], row["city"]))
            set_travel_detail(
                card,
                soup,
                "Trajet hôtel ↔ aéroport",
                "Trajet à vérifier sur Google Maps",
                google_maps_directions(row["name"], row["city"]),
            )
        if row.get("official_url"):
            link = card.select_one(".hotel-link")
            if link:
                link["href"] = row["official_url"]

    recalculate_counts(soup)
    backup = write_soup(catalogue_path, soup)
    return {
        "rows": len(rows),
        "updated": updated,
        "created": created,
        "missing_city_sections": sorted(set(filter(None, missing_city_sections))),
        "backup": str(backup),
    }


ADMIN_HTML = """<!doctype html>
<html lang="fr">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>STOURS Catalogue Admin</title>
  <style>
    :root{--blue:#1628A9;--red:#A8371D;--warm:#FFFEE9;--ink:#141414;--line:#e8dfc6}
    body{margin:0;font-family:Inter,Arial,sans-serif;background:var(--warm);color:var(--ink)}
    header{position:sticky;top:0;background:#fffef1;border-bottom:1px solid var(--line);z-index:2;padding:16px 24px;display:flex;gap:16px;align-items:center}
    h1{font-family:Georgia,serif;color:var(--blue);margin:0;font-size:28px}.small{font-size:12px;color:#6d6658}
    main{display:grid;grid-template-columns:360px 1fr;gap:18px;padding:18px}
    aside,.panel{background:white;border:1px solid var(--line);border-radius:18px;box-shadow:0 18px 50px rgba(20,20,20,.08)}
    aside{height:calc(100vh - 104px);overflow:auto}.panel{padding:20px}
    .toolbar{display:flex;gap:8px;padding:12px;border-bottom:1px solid var(--line);position:sticky;top:0;background:white;border-radius:18px 18px 0 0}
    input,textarea,select{width:100%;box-sizing:border-box;border:1px solid var(--line);border-radius:10px;padding:10px;background:#fff;font:inherit}
    textarea{min-height:84px;resize:vertical}button{border:0;background:var(--blue);color:white;border-radius:999px;padding:10px 14px;font-weight:800;cursor:pointer}
    button.secondary{background:var(--red)}button.light{background:#f4edda;color:var(--blue)}
    .hotel{padding:12px 14px;border-bottom:1px solid var(--line);cursor:pointer}.hotel:hover,.hotel.active{background:#fff8d6}.hotel b{display:block;color:var(--blue)}
    label{display:block;font-size:12px;text-transform:uppercase;letter-spacing:.08em;font-weight:900;color:var(--red);margin:14px 0 6px}
    .grid{display:grid;grid-template-columns:1fr 1fr;gap:14px}.images{display:grid;grid-template-columns:repeat(auto-fill,minmax(140px,1fr));gap:12px}
    .image-card{border:1px dashed var(--red);border-radius:14px;padding:8px}.image-card img{width:100%;height:95px;object-fit:cover;border-radius:10px;background:#eee}
    .status{white-space:pre-wrap;background:#141414;color:#fff;border-radius:12px;padding:12px;margin-top:12px;min-height:42px}
    @media(max-width:900px){main{grid-template-columns:1fr}aside{height:360px}}
  </style>
</head>
<body>
<header>
  <div>
    <h1>STOURS Catalogue Admin</h1>
    <div class="small">Modifier contenu, télécharger photos localement, switcher images, importer noms depuis Excel.</div>
  </div>
  <button onclick="window.open('/', '_blank')" class="light">Voir catalogue</button>
</header>
<main>
  <aside>
    <div class="toolbar"><input id="search" placeholder="Chercher hôtel..." oninput="renderList()"></div>
    <div id="list"></div>
  </aside>
  <section class="panel">
    <div class="grid">
      <div><label>Importer Excel (.xlsx)</label><input id="excelPath" placeholder="/chemin/fichier.xlsx"><button onclick="previewExcel()" class="light">Prévisualiser</button> <button onclick="applyExcel()" class="secondary">Appliquer noms/liens</button></div>
      <div><label>Résultat</label><div id="status" class="status">Chargement...</div></div>
    </div>
    <hr>
    <div id="editor">Sélectionnez un hôtel.</div>
  </section>
</main>
<script>
let hotels = [];
let current = null;
const labels = ['ROOMS','RESTAURANTS','SIGNATURE EXPERIENCE','FACILITIES'];
const $ = id => document.getElementById(id);
function status(value){ $('status').textContent = typeof value === 'string' ? value : JSON.stringify(value,null,2); }
async function api(path, options={}){
  const res = await fetch(path, {headers:{'Content-Type':'application/json'}, ...options});
  const data = await res.json();
  if(!res.ok) throw new Error(data.error || res.statusText);
  return data;
}
async function loadHotels(){
  hotels = (await api('/api/hotels')).hotels;
  renderList();
  status(`${hotels.length} hôtels chargés.`);
}
function renderList(){
  const q = $('search').value.toLowerCase();
  $('list').innerHTML = hotels.filter(h => !q || h.name.toLowerCase().includes(q) || h.city.name.toLowerCase().includes(q)).map(h =>
    `<div class="hotel ${current && current.index===h.index?'active':''}" onclick="edit(${h.index})"><b>${h.name}</b><span>${h.city.name}</span></div>`
  ).join('');
}
function edit(index){
  current = hotels.find(h => h.index === index);
  renderList();
  const s = current.sections;
  $('editor').innerHTML = `
    <div class="grid">
      <div><label>Nom hôtel</label><input id="name" value="${escapeHtml(current.name)}"></div>
      <div><label>Classification</label><input id="classification" value="${escapeHtml(current.classification)}"></div>
      <div><label>Localisation Google Maps</label><input id="location" value="${escapeHtml(current.location)}"></div>
      <div><label>Trajet hôtel ↔ aéroport</label><input id="airport_route" value="${escapeHtml(current.airport_route)}"></div>
      <div><label>Noms/types de chambres</label><input id="room_names" value="${escapeHtml(current.room_names)}"></div>
      <div><label>Site officiel</label><input id="official_url" value="${escapeHtml(current.official_url)}"></div>
    </div>
    ${labels.map(label => `<label>${label}</label><textarea id="section-${label}">${escapeHtml(s[label] || '')}</textarea>`).join('')}
    <label>Photos locales (slot 1 = photo principale)</label>
    <div class="images">${current.images.map((src,i)=>imageCard(src,i)).join('')}${imageCard('', current.images.length)}</div>
    <p><button onclick="saveHotel()">Enregistrer contenu/photos</button></p>
  `;
}
function imageCard(src, i){
  return `<div class="image-card">
    ${src ? `<img src="/${src}">` : `<div style="height:95px;display:grid;place-items:center;background:#f3efe0;border-radius:10px">Nouveau slot</div>`}
    <label>Slot ${i+1}</label><input id="image-${i}" value="${escapeHtml(src)}" placeholder="assets/...jpg">
    <input id="download-${i}" placeholder="URL image à télécharger">
    <button class="light" onclick="downloadPhoto(${i})">Télécharger ici</button>
  </div>`;
}
function escapeHtml(value){return String(value||'').replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));}
async function saveHotel(){
  if(!current) return;
  const images = [];
  for(let i=0;i<12;i++){ const el = $('image-'+i); if(el && el.value.trim()) images.push(el.value.trim()); }
  const payload = {
    name:$('name').value, classification:$('classification').value, location:$('location').value,
    airport_route:$('airport_route').value, room_names:$('room_names').value, official_url:$('official_url').value, images,
    sections:Object.fromEntries(labels.map(label => [label, $('section-'+label).value]))
  };
  const data = await api(`/api/hotels/${current.index}`, {method:'POST', body:JSON.stringify(payload)});
  status(data);
  await loadHotels();
  edit(data.hotel.index);
}
async function downloadPhoto(slot){
  const url = $('download-'+slot).value.trim();
  if(!url) return status('Ajoutez une URL image.');
  const data = await api(`/api/hotels/${current.index}/download-photo`, {method:'POST', body:JSON.stringify({url, slot})});
  status(data);
  await loadHotels();
  edit(data.hotel.index);
}
async function previewExcel(){
  const data = await api('/api/import-excel', {method:'POST', body:JSON.stringify({path:$('excelPath').value, apply:false})});
  status(data);
}
async function applyExcel(){
  const data = await api('/api/import-excel', {method:'POST', body:JSON.stringify({path:$('excelPath').value, apply:true})});
  status(data);
  await loadHotels();
}
loadHotels().catch(err => status(err.message));
</script>
</body>
</html>"""


class AdminHandler(SimpleHTTPRequestHandler):
    catalogue_path = DEFAULT_CATALOGUE

    def translate_path(self, path: str) -> str:
        requested = unquote(urlparse(path).path).lstrip("/")
        return str((ROOT / requested).resolve())

    def send_json(self, payload: object, status: HTTPStatus = HTTPStatus.OK) -> None:
        body = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_json(self) -> dict[str, object]:
        length = int(self.headers.get("content-length", "0"))
        if not length:
            return {}
        return json.loads(self.rfile.read(length).decode("utf-8"))

    def do_GET(self) -> None:
        path = urlparse(self.path).path
        if path == "/admin":
            body = ADMIN_HTML.encode("utf-8")
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if path == "/api/hotels":
            self.send_json({"hotels": all_hotels(self.catalogue_path)})
            return
        super().do_GET()

    def do_POST(self) -> None:
        try:
            path = urlparse(self.path).path
            payload = self.read_json()
            match_update = re.fullmatch(r"/api/hotels/(\d+)", path)
            match_photo = re.fullmatch(r"/api/hotels/(\d+)/download-photo", path)
            if match_update:
                self.send_json(update_hotel(self.catalogue_path, int(match_update.group(1)), payload))
                return
            if match_photo:
                self.send_json(
                    download_photo(
                        self.catalogue_path,
                        int(match_photo.group(1)),
                        str(payload.get("url", "")),
                        int(payload.get("slot", 0)),
                    )
                )
                return
            if path == "/api/import-excel":
                excel_path = Path(str(payload.get("path", ""))).expanduser()
                if not excel_path.is_absolute():
                    excel_path = ROOT / excel_path
                if not excel_path.exists():
                    raise ValueError(f"Excel file not found: {excel_path}")
                if payload.get("apply"):
                    self.send_json(apply_excel_import(self.catalogue_path, excel_path))
                else:
                    rows = parse_excel(excel_path)
                    self.send_json({"rows": len(rows), "hotels": rows[:200]})
                return
            self.send_json({"error": "Unknown endpoint"}, HTTPStatus.NOT_FOUND)
        except Exception as exc:
            self.send_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the local STOURS catalogue admin app.")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8000)
    parser.add_argument("--catalogue", type=Path, default=DEFAULT_CATALOGUE)
    args = parser.parse_args()
    AdminHandler.catalogue_path = args.catalogue.resolve()
    server = ThreadingHTTPServer((args.host, args.port), AdminHandler)
    print(f"STOURS admin app: http://{args.host}:{args.port}/admin")
    print(f"Catalogue file: {AdminHandler.catalogue_path}")
    server.serve_forever()


if __name__ == "__main__":
    main()
